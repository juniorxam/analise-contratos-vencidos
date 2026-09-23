import re
import io
import csv
import zipfile
import unicodedata
from xml.sax.saxutils import escape
import streamlit as st
import pandas as pd
import numpy as np
from datetime import datetime
from io import BytesIO
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

try:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER, TA_LEFT
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import (
        SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak
    )
    HAS_REPORTLAB = True
except ImportError:
    HAS_REPORTLAB = False

try:
    import chardet
    HAS_CHARDET = True
except ImportError:
    HAS_CHARDET = False

st.set_page_config(
    page_title="Análise de Contratos Vencidos",
    page_icon="📄",
    layout="wide",
    initial_sidebar_state="expanded",
)

st.markdown("""
<style>
    :root { --navy:#16324F; --teal:#0F766E; --mint:#ECFDF5; --ink:#172033; --muted:#64748B; }
    .stApp { background: linear-gradient(180deg, #F8FAFC 0%, #FFFFFF 38%); color: var(--ink); }
    [data-testid="stSidebar"] { background: linear-gradient(180deg, #102A43 0%, #16324F 100%); }
    [data-testid="stSidebar"] * { color: #F8FAFC !important; }
    [data-testid="stSidebar"] hr { border-color: rgba(255,255,255,.2); }
    .hero { padding: 26px 30px; border-radius: 20px; margin: 4px 0 24px;
        background: radial-gradient(circle at 90% 10%, #2DD4BF 0, transparent 30%),
                    linear-gradient(120deg, #16324F 0%, #0F766E 100%);
        color: white; box-shadow: 0 12px 28px rgba(15,118,110,.18); }
    .hero h1 { margin: 0; font-size: 2rem; letter-spacing: -.03em; }
    .hero p { margin: 8px 0 0; color: #D9FDF5; font-size: .98rem; }
    div[data-testid="stMetric"] { background: white; border: 1px solid #E2E8F0;
        border-radius: 14px; padding: 12px 14px; box-shadow: 0 5px 16px rgba(15,23,42,.05); }
    div[data-testid="stMetricLabel"] { color: #64748B; }
    div[data-testid="stMetricValue"] { color: #16324F; }
    .section-kicker { text-transform: uppercase; letter-spacing: .12em; color: #0F766E;
        font-size: .72rem; font-weight: 700; margin-bottom: 4px; }
    .stTabs [data-baseweb="tab-list"] { gap: 8px; border-bottom: 1px solid #E2E8F0; }
    .stTabs [data-baseweb="tab"] { border-radius: 10px 10px 0 0; padding: 11px 15px; }
    .stTabs [aria-selected="true"] { background: #ECFDF5; color: #0F766E; font-weight: 700; }
    .stButton button, .stDownloadButton button { border-radius: 10px; font-weight: 650; }
</style>
""", unsafe_allow_html=True)

EXPECTED_COLS = [
    "ESCALA", "DESC. ESCALA", "NUMFUNC", "NUMVINC", "SERVIDOR", "CPF",
    "SETOR", "VINCULO", "DATA DE INICIO - VINCULO", "CARGO", "OCUPACAO",
    "CARGA HORARIA ESCALADA", "CARGA HORARIA",
]

ACCEPTED_EXT = ["csv", "xls", "xlsx", "xlsm"]


def format_br_number(value, decimals=2, signed=False):
    """Formata número no padrão brasileiro: 1.234,56."""
    if value is None or (isinstance(value, float) and np.isnan(value)):
        return ""
    try:
        value = float(value)
        raw = f"{value:+,.{decimals}f}" if signed else f"{value:,.{decimals}f}"
        return raw.replace(",", "_").replace(".", ",").replace("_", ".")
    except (TypeError, ValueError):
        return str(value)


def format_br_int(value):
    if value is None:
        return ""
    try:
        return f"{int(value):,}".replace(",", ".")
    except (TypeError, ValueError):
        return str(value)


def format_dataframe_for_display(data):
    """Cria uma cópia amigável sem alterar os dados usados nos cálculos."""
    out = data.copy()
    for col in out.columns:
        if pd.api.types.is_datetime64_any_dtype(out[col]):
            out[col] = out[col].dt.strftime("%d/%m/%Y")
        elif pd.api.types.is_numeric_dtype(out[col]):
            is_count = any(token in str(col).upper() for token in (
                "QTD", "QUANT", "ALOC", "PESSOAS", "SERVIDORES", "DIGITOS", "REGISTROS"
            ))
            out[col] = out[col].map(
                lambda x: format_br_int(x) if is_count else format_br_number(x)
            )
    return out


def display_dataframe(data, *args, **kwargs):
    st.dataframe(format_dataframe_for_display(data), *args, **kwargs)


# ==================================================================
# LEITURA ROBUSTA — CSV / XLS / XLSX
# ==================================================================
def sniff_file_type(file_bytes: bytes) -> str:
    head = file_bytes[:8]
    if head[:4] == b"PK\x03\x04":
        return "xlsx"
    if head == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1":
        return "xls"
    return "text"


def detect_encoding(file_bytes: bytes) -> str:
    if HAS_CHARDET:
        enc = chardet.detect(file_bytes[:100_000]).get("encoding")
        if enc:
            return enc
    for enc in ("utf-8-sig", "utf-8", "latin1", "cp1252"):
        try:
            file_bytes.decode(enc)
            return enc
        except Exception:
            continue
    return "latin1"


def read_text_robust(file_bytes: bytes) -> pd.DataFrame:
    """Lê CSV respeitando aspas, delimitadores em textos e BOM."""
    enc = detect_encoding(file_bytes)
    text = file_bytes.decode(enc, errors="replace").lstrip("\ufeff")
    if not text.strip():
        return pd.DataFrame()
    sample = text[:10000]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters="\t;,|")
        sep = dialect.delimiter
    except csv.Error:
        first = next((ln for ln in text.splitlines() if ln.strip()), "")
        counts = {d: first.count(d) for d in ["\t", ";", "|", ","]}
        sep = max(counts, key=counts.get) if max(counts.values()) else ","
    return pd.read_csv(
        io.StringIO(text), sep=sep, engine="python", header=None,
        dtype=object, keep_default_na=False, on_bad_lines="warn"
    )


def read_excel_robust(file_bytes: bytes, kind: str) -> pd.DataFrame:
    bio = BytesIO(file_bytes)
    engine = "xlrd" if kind == "xls" else "openpyxl"
    try:
        return pd.read_excel(bio, engine=engine, header=None, dtype=object)
    except Exception:
        bio.seek(0)
        return pd.read_excel(bio, header=None, dtype=object)


@st.cache_data(show_spinner=False)
def read_any(file_bytes: bytes, file_name: str) -> pd.DataFrame:
    kind = sniff_file_type(file_bytes)
    if file_name.lower().endswith(".csv"):
        return read_text_robust(file_bytes)
    if kind == "xls":
        return read_excel_robust(file_bytes, "xls")
    if kind == "xlsx":
        return read_excel_robust(file_bytes, "xlsx")
    return read_text_robust(file_bytes)


# ==================================================================
# PARSE
# ==================================================================
def find_header_index(col_series):
    for i in range(min(15, len(col_series))):
        v = str(col_series.iloc[i])
        if "NUMFUNC" in v and "ESCALA" in v:
            return i
    return None


def parse_raw(raw):
    raw = raw.dropna(how="all").reset_index(drop=True)

    if raw.shape[1] <= 3:
        col = raw.iloc[:, 0].fillna("").astype(str)
        hi = find_header_index(col)
        if hi is None:
            hi = 0
        header = [h.strip() for h in col.iloc[hi].split("\t")]
        body = col.iloc[hi + 1:].str.split("\t", expand=True)
        n = min(body.shape[1], len(header))
        body = body.iloc[:, :n]
        body.columns = header[:n]
        return body.reset_index(drop=True)

    hi = 0
    for i in range(min(15, len(raw))):
        vals = [str(v) for v in raw.iloc[i].tolist()]
        if any("NUMFUNC" in v for v in vals):
            hi = i
            break
    df = raw.iloc[hi + 1:].copy()
    df.columns = [str(v).strip() for v in raw.iloc[hi].tolist()]
    return df.reset_index(drop=True)


# ==================================================================
# HELPERS DE NORMALIZAÇÃO / CPF
# ==================================================================
def safe_str(v):
    if v is None:
        return ""
    try:
        if pd.isna(v):
            return ""
    except (TypeError, ValueError):
        pass
    s = str(v).strip()
    return "" if s.lower() in ("nan", "none", "nat") else s


def to_int_str(v):
    s = safe_str(v)
    if s == "":
        return ""
    if re.fullmatch(r"-?\d+\.0+", s):
        return s.split(".")[0]
    return s


def only_digits(v: str) -> str:
    return re.sub(r"\D", "", str(v))


def normalize_identifier(v) -> str:
    s = safe_str(v)
    if re.fullmatch(r"\d+\.0+", s):
        s = s.split(".", 1)[0]
    return only_digits(s)


def cpf_is_valid(cpf: str) -> bool:
    c = only_digits(cpf)
    if len(c) != 11 or len(set(c)) == 1:
        return False
    total = sum(int(c[i]) * (10 - i) for i in range(9))
    d1 = (total * 10 % 11) % 10
    total = sum(int(c[i]) * (11 - i) for i in range(10))
    d2 = (total * 10 % 11) % 10
    return c[9] == str(d1) and c[10] == str(d2)


def format_cpf(cpf: str) -> str:
    c = only_digits(cpf)
    if len(c) == 11:
        return f"{c[:3]}.{c[3:6]}.{c[6:9]}-{c[9:]}"
    return c


def strip_leading_zeros(s: str) -> str:
    s = only_digits(s)
    if s == "":
        return ""
    stripped = s.lstrip("0")
    return stripped if stripped != "" else "0"


def sanitize_filename(name: str) -> str:
    """Remove caracteres inválidos para nome de arquivo."""
    name = str(name).strip()
    name = re.sub(r'[\\/:*?"<>|\r\n\t]+', "_", name)
    name = re.sub(r"\s+", " ", name).strip(" .")
    return name[:80] if name else "SEM_SETOR"


def normalize_scale_label(value) -> str:
    """Corrige escapes indevidos e espaços nos nomes de escalas."""
    text = safe_str(value)
    # Corrige tanto o texto literal '/t' e '\\t' quanto variantes com espaços.
    text = re.sub(r"[/\\]\s*t\b", " ", text, flags=re.IGNORECASE)
    text = text.replace("\t", " ")
    text = re.sub(r"\s+", " ", text).strip()
    return text


CARGO_CANONICOS = {
    "medico rqe": "Médico - RQE",
    "farmaceutico bioquimico": "Farmacêutico-Bioquímico",
    "tecnico em imobilizacao ortopedica": "Técnico em Imobilização Ortopédica",
    "tecnico eletricista": "Técnico Eletricista",
    "auxilar em servicos de saude i": "Auxiliar em Serviços de Saúde I",
    "auxilar em servicos de saude ii": "Auxiliar em Serviços de Saúde II",
    "auxiliar em servicos de saude i": "Auxiliar em Serviços de Saúde I",
    "auxiliar em servicos de saude ii": "Auxiliar em Serviços de Saúde II",
    "assistente em servicos de saude i": "Assistente em Serviços de Saúde I",
    "assistente em servicos de saude ii": "Assistente em Serviços de Saúde II",
    "auxiliar de servicos hospitalares": "Auxiliar de Serviços Hospitalares",
    "pesquisador docente em saude publica": "Pesquisador Docente em Saúde Pública",
    "analista tecnico administrativo": "Analista Técnico-Administrativo",
    "diretor de integracao multiprofissional hospitalar concentracao ii das": "Diretor de Integração Multiprofissional Hospitalar - Concentração II DAS",
    "diretor administrativo de unidade porte dai 1": "Diretor Administrativo de Unidade Porte DAI-1",
    "motorista de ambulancia": "Motorista de Ambulância",
    "motorista de representacao": "Motorista de Representação",
    "analista em tecnologia da informacao": "Analista em Tecnologia da Informação",
    "analista em controle de zoonoses": "Analista em Controle de Zoonoses",
    "diretor geral de unidade porte 1 dai": "Diretor Geral de Unidade Porte 1 DAI",
    "diretor geral de unidade porte 2 das": "Diretor Geral de Unidade Porte 2 DAS",
    "diretor geral de unidade porte 3 das": "Diretor Geral de Unidade Porte 3 DAS",
    "tecnico em enfermagem": "Técnico em Enfermagem",
    "tecnica em enfermagem": "Técnica em Enfermagem",
    "auxiliar de enfermagem": "Auxiliar de Enfermagem",
    "tecnico de enfermagem": "Técnico de Enfermagem",
    "tecnica de enfermagem": "Técnica de Enfermagem",
    "enfermeiro obstetra": "Enfermeiro Obstetra",
    "medico clinico": "Médico Clínico",
    "medico plantonista": "Médico Plantonista",
    "medico": "Médico", "medica": "Médica",
    "tecnico": "Técnico", "tecnica": "Técnica",
    "enfermeiro": "Enfermeiro", "enfermeira": "Enfermeira",
    "farmaceutico": "Farmacêutico", "farmaceutica": "Farmacêutica",
    "fisioterapeuta": "Fisioterapeuta", "fonoaudiologo": "Fonoaudiólogo",
    "psicologo": "Psicólogo", "psicologa": "Psicóloga",
    "odontologo": "Odontólogo", "odontologa": "Odontóloga",
    "biomedico": "Biomédico", "biomedica": "Biomédica",
    "nutricionista": "Nutricionista", "assistente social": "Assistente Social",
    "agente comunitario de saude": "Agente Comunitário de Saúde",
    "cirurgiao dentista": "Cirurgião-Dentista",
    "terapeuta ocupacional": "Terapeuta Ocupacional",
    "educador fisico": "Educador Físico", "educadora fisica": "Educadora Física",
    "veterinario": "Veterinário", "veterinaria": "Veterinária",
    "engenheiro": "Engenheiro", "engenheira": "Engenheira",
    "administrador": "Administrador", "administradora": "Administradora",
    "coordenador": "Coordenador", "coordenadora": "Coordenadora",
    "supervisor": "Supervisor", "supervisora": "Supervisora",
    "auxiliar": "Auxiliar", "recepcionista": "Recepcionista",
    "secretario": "Secretário", "secretaria": "Secretária",
}

CARGO_TOKENS = {
    "assistente": "Assistente", "servicos": "Serviços", "saude": "Saúde",
    "cirurgiao": "Cirurgião", "dentista": "Dentista", "laboratorio": "Laboratório",
    "motorista": "Motorista", "ambulancia": "Ambulância", "bioquimico": "Bioquímico",
    "analista": "Analista", "desenvolvimento": "Desenvolvimento", "social": "Social",
    "administrador": "Administrador", "hospitalar": "Hospitalar", "assessor": "Assessor",
    "comissionado": "Comissionado", "apoio": "Apoio", "executivo": "Executivo",
    "gestor": "Gestor", "seguranca": "Segurança", "trabalho": "Trabalho",
    "professor": "Professor", "educacao": "Educação", "basica": "Básica",
    "maqueiro": "Maqueiro", "radiologia": "Radiologia", "clinico": "Clínico",
    "instrumentador": "Instrumentador", "cirurgico": "Cirúrgico", "controle": "Controle",
    "zoonoses": "Zoonoses", "informacao": "Informação", "contabilidade": "Contabilidade",
    "microcomputador": "Microcomputador", "gerais": "Gerais", "geral": "Geral",
    "porte": "Porte", "unidade": "Unidade", "diretor": "Diretor",
    "integracao": "Integração", "multiprofissional": "Multiprofissional",
    "concentracao": "Concentração", "administrativo": "Administrativo",
    "financeiro": "Financeiro", "representacao": "Representação", "economista": "Economista",
    "publica": "Pública", "pesquisador": "Pesquisador", "docente": "Docente",
    "eletricista": "Eletricista", "imobilizacao": "Imobilização", "ortopedica": "Ortopédica",
    "biblioteconomista": "Biblioteconomista", "museologo": "Museólogo", "pedagogo": "Pedagogo",
    "organizacional": "Organizacional", "enfermagem": "Enfermagem", "fisioterapeuta": "Fisioterapeuta",
    "nutricionista": "Nutricionista", "terapeuta": "Terapeuta", "ocupacional": "Ocupacional",
    "operador": "Operador", "especializado": "Especializado",
    "tecnico": "Técnico", "tecnica": "Técnica",
    "medico": "Médico", "medica": "Médica",
    "enfermeiro": "Enfermeiro", "enfermeira": "Enfermeira",
    "farmaceutico": "Farmacêutico", "farmaceutica": "Farmacêutica",
    "psicologo": "Psicólogo", "psicologa": "Psicóloga",
    "fonoaudiologo": "Fonoaudiólogo", "fonoaudiologa": "Fonoaudióloga",
    "odontologo": "Odontólogo", "odontologa": "Odontóloga",
    "biomedico": "Biomédico", "biomedica": "Biomédica",
    "comunitario": "Comunitário", "comunitaria": "Comunitária",
    "saude": "Saúde", "fisico": "Físico", "fisica": "Física",
    "cirurgiao": "Cirurgião", "cirurgia": "Cirurgia",
}

CARGO_TOKENS.update({
    "auxilar": "Auxiliar", "auxiliar": "Auxiliar", "secretario": "Secretário",
    "secretaria": "Secretária", "i": "I", "ii": "II", "iii": "III", "iv": "IV",
    "v": "V", "rqe": "RQE", "dai": "DAI", "das": "DAS", "cea": "CEA",
    "cca": "CCA", "ae": "AE", "ca": "CA",
})


def cargo_key(value) -> str:
    text = re.sub(r"\s+", " ", safe_str(value)).strip().casefold()
    text = "".join(
        char for char in unicodedata.normalize("NFKD", text)
        if not unicodedata.combining(char)
    )
    return re.sub(r"[^a-z0-9]+", " ", text).strip()


def title_case_cargo(value) -> str:
    text = re.sub(r"\s+", " ", safe_str(value)).strip()
    if not text:
        return ""
    if text.isupper() and len(text) <= 6:
        return text
    words = text.lower().split(" ")
    connectors = {"de", "da", "do", "das", "dos", "e", "em"}
    return " ".join(
        word if i > 0 and word in connectors else word.capitalize()
        for i, word in enumerate(words)
    )


def standardize_cargo_column(series):
    """Padroniza cargos e retorna (valor padronizado, status da correção)."""
    originals = series.apply(safe_str)
    def standardize(value):
        key = cargo_key(value)
        if key in CARGO_CANONICOS:
            return CARGO_CANONICOS[key]
        words = key.split()
        if not words:
            return ""
        connectors = {"de", "da", "do", "das", "dos", "e", "em"}
        return " ".join(
            word if word in connectors else CARGO_TOKENS.get(word, word.capitalize())
            for word in words
        )
    standardized = originals.map(standardize)
    status = np.where(
        (originals != "") & (originals != standardized), "CORRIGIDO", "OK"
    )
    return standardized, pd.Series(status, index=series.index)


# ==================================================================
# LIMPEZA
# ==================================================================
def clean(df):
    df = df.copy()
    for c in EXPECTED_COLS:
        if c not in df.columns:
            df[c] = ""
    df = df[EXPECTED_COLS]

    for c in df.columns:
        df[c] = df[c].apply(safe_str)

    for c in ["ESCALA", "DESC. ESCALA"]:
        df[c] = df[c].apply(normalize_scale_label)

    df["CARGO_ORIGINAL"] = df["CARGO"]
    df["CARGO"], df["CARGO_STATUS"] = standardize_cargo_column(df["CARGO"])
    df["OCUPACAO_ORIGINAL"] = df["OCUPACAO"]
    df["OCUPACAO"], df["OCUPACAO_STATUS"] = standardize_cargo_column(df["OCUPACAO"])

    df["NUMFUNC"] = df["NUMFUNC"].apply(to_int_str)
    df["NUMVINC"] = df["NUMVINC"].apply(to_int_str)

    df["CARGA HORARIA ESCALADA"] = pd.to_numeric(
        df["CARGA HORARIA ESCALADA"].str.replace(",", ".", regex=False),
        errors="coerce",
    )
    df["CARGA HORARIA"] = pd.to_numeric(
        df["CARGA HORARIA"].str.replace(",", ".", regex=False),
        errors="coerce",
    )

    df["CPF"] = df["CPF"].apply(normalize_identifier)
    df["CPF_FORMATADO"] = df["CPF"].apply(format_cpf)
    df["CPF_QTD_DIGITOS"] = df["CPF"].str.len()
    df["CPF_VALIDO"] = df["CPF"].apply(cpf_is_valid)

    def status_cpf(row):
        n = row["CPF_QTD_DIGITOS"]
        if n == 0:
            return "AUSENTE"
        if n != 11:
            return "INCOMPLETO" if n < 11 else "EXCEDENTE"
        return "OK" if row["CPF_VALIDO"] else "DIGITO_VERIFICADOR_INVALIDO"
    df["CPF_STATUS"] = df.apply(status_cpf, axis=1)

    nf_norm = df["NUMFUNC"].apply(strip_leading_zeros)
    cp_norm = df["CPF"].apply(strip_leading_zeros)
    df["_NF_NORM"] = nf_norm
    df["_CP_NORM"] = cp_norm
    df["ALERTA_NUMFUNC_CPF"] = np.where(
        (df["_NF_NORM"] != "")
        & (df["_CP_NORM"] != "")
        & (df["_NF_NORM"] != "0")
        & (df["_NF_NORM"] == df["_CP_NORM"])
        & (df["_NF_NORM"].str.len() >= 8),
        "S", "",
    )
    df = df.drop(columns=["_NF_NORM", "_CP_NORM"])

    df["DATA DE INICIO - VINCULO"] = pd.to_datetime(
        df["DATA DE INICIO - VINCULO"], errors="coerce", dayfirst=True
    )

    df["CHAVE_ID"] = np.where(
        (df["NUMFUNC"] != "") & (df["NUMVINC"] != ""),
        df["NUMFUNC"] + "-" + df["NUMVINC"],
        "",
    )
    df["PESSOA_ID"] = np.where(
        df["CPF_VALIDO"], "CPF:" + df["CPF"],
        np.where(df["NUMFUNC"] != "", "NF:" + df["NUMFUNC"], "")
    )
    df["ESCALA_ID"] = np.where(
        (df["ESCALA"] != "") | (df["DESC. ESCALA"] != ""),
        df["ESCALA"] + " — " + df["DESC. ESCALA"], "(SEM ESCALA)"
    )

    df = df[
        (df["NUMFUNC"] != "") | (df["NUMVINC"] != "")
        | (df["SERVIDOR"] != "") | (df["CPF"] != "")
    ].reset_index(drop=True)

    base_order = [
        "CHAVE_ID", "PESSOA_ID", "ESCALA_ID", "NUMFUNC", "NUMVINC", "SERVIDOR",
        "CPF", "CPF_FORMATADO", "CPF_QTD_DIGITOS", "CPF_STATUS", "CPF_VALIDO",
        "ALERTA_NUMFUNC_CPF",
        "SETOR", "VINCULO", "DATA DE INICIO - VINCULO",
        "CARGO", "OCUPACAO",
        "CARGA HORARIA ESCALADA", "CARGA HORARIA",
        "ESCALA", "DESC. ESCALA",
    ]
    base_order = [c for c in base_order if c in df.columns]
    rest = [c for c in df.columns if c not in base_order]
    return df[base_order + rest]


# ==================================================================
# ALERTAS
# ==================================================================
def alert_cpf_incompleto(df):
    mask = (df["CPF_QTD_DIGITOS"] > 0) & (df["CPF_STATUS"] != "OK")
    out = df.loc[mask].copy()
    if len(out) == 0:
        return pd.DataFrame(columns=[
            "CHAVE_ID", "NUMFUNC", "NUMVINC", "SERVIDOR",
            "CPF", "CPF_FORMATADO", "CPF_QTD_DIGITOS", "CPF_STATUS",
            "SUGESTAO_CPF", "PROBLEMA",
            "SETOR", "VINCULO", "CARGO", "OCUPACAO",
        ])

    def sugerir(c):
        n = len(c)
        if 0 < n < 11:
            return c.zfill(11)
        return ""

    out["SUGESTAO_CPF"] = out["CPF"].apply(sugerir)
    out["PROBLEMA"] = out.apply(
        lambda r: (
            f"CPF com {r['CPF_QTD_DIGITOS']} dígito(s) — o esperado são 11"
            if r["CPF_STATUS"] != "DIGITO_VERIFICADOR_INVALIDO"
            else "CPF possui 11 dígitos, mas os dígitos verificadores são inválidos"
        ), axis=1
    )
    cols = [
        "CHAVE_ID", "NUMFUNC", "NUMVINC", "SERVIDOR",
        "CPF", "CPF_FORMATADO", "CPF_QTD_DIGITOS", "CPF_STATUS",
        "SUGESTAO_CPF", "PROBLEMA",
        "SETOR", "VINCULO", "CARGO", "OCUPACAO",
    ]
    cols = [c for c in cols if c in out.columns]
    return out[cols].sort_values(
        ["CPF_QTD_DIGITOS", "SERVIDOR"], ascending=[True, True]
    )


def alert_numfunc_cpf(df):
    out = df.loc[df["ALERTA_NUMFUNC_CPF"] == "S"].copy()
    if len(out) == 0:
        return out
    out["PROBLEMA"] = (
        "NUMFUNC contém o CPF — corrigir número funcional "
        "(comparação tolera zeros à esquerda)"
    )
    return out


def alert_contratos_vencidos(df, ref_date):
    d = df.copy()
    d["DIAS_DESDE_INICIO"] = (
        pd.Timestamp(ref_date) - d["DATA DE INICIO - VINCULO"]
    ).dt.days
    mask_temp = d["VINCULO"].str.contains("Tempor", case=False, na=False)
    limite = pd.Timestamp(ref_date) - pd.DateOffset(years=2)
    mask_venc = d["DATA DE INICIO - VINCULO"] < limite
    out = d.loc[mask_temp & mask_venc].copy()
    out["ANOS_DESDE_INICIO"] = (out["DIAS_DESDE_INICIO"] / 365.25).round(2)
    cols = [
        "CHAVE_ID", "NUMFUNC", "NUMVINC", "SERVIDOR",
        "CPF_FORMATADO", "SETOR", "VINCULO",
        "DATA DE INICIO - VINCULO", "ANOS_DESDE_INICIO",
        "CARGO", "OCUPACAO",
        "CARGA HORARIA ESCALADA", "CARGA HORARIA",
    ]
    cols = [c for c in cols if c in out.columns]
    return out[cols].sort_values("ANOS_DESDE_INICIO", ascending=False)


# ==================================================================
# AGREGAÇÕES
# ==================================================================
def build_cbo(df):
    if len(df) == 0:
        return pd.DataFrame(columns=[
            "CARGO", "OCUPACAO", "QTD_ALOCACOES",
            "QTD_SERVIDORES_UNICOS",
            "CH_ESCALADA_TOTAL", "CH_CONTRATADA_TOTAL",
        ])
    g = df.groupby(["CARGO", "OCUPACAO"], dropna=False).agg(
        QTD_ALOCACOES=("CHAVE_ID", "size"),
        QTD_SERVIDORES_UNICOS=("PESSOA_ID", lambda x: x.replace("", pd.NA).dropna().nunique()),
        CH_ESCALADA_TOTAL=("CARGA HORARIA ESCALADA", "sum"),
        CH_CONTRATADA_TOTAL=("CARGA HORARIA", "sum"),
    ).reset_index()
    return g.sort_values(
        "QTD_SERVIDORES_UNICOS", ascending=False
    ).reset_index(drop=True)


def build_resumo(df, alerta_num, alerta_cpf, alerta_temp):
    if len(df) == 0:
        return pd.DataFrame()
    resumo = {
        "Total de alocações (linhas)": len(df),
        "Pessoas únicas":
            df["PESSOA_ID"].replace("", np.nan).dropna().nunique(),
        "Vínculos únicos":
            df["CHAVE_ID"].replace("", np.nan).dropna().nunique(),
        "CPFs válidos (dígitos verificadores)":
            int(df["CPF_VALIDO"].sum()),
        "CPFs únicos válidos":
            df.loc[df["CPF_VALIDO"], "CPF"].nunique(),
        "CPFs incompletos / excedentes": len(alerta_cpf),
        "Registros com NUMFUNC = CPF": len(alerta_num),
        "Contratos temporários vencidos (>2a)": len(alerta_temp),
        "Setores distintos":
            df["SETOR"].replace("", np.nan).dropna().nunique(),
        "Cargos distintos":
            df["CARGO"].replace("", np.nan).dropna().nunique(),
        "Cargos corrigidos": int((df["CARGO_STATUS"] == "CORRIGIDO").sum()),
        "Ocupações (CBO) distintas":
            df["OCUPACAO"].replace("", np.nan).dropna().nunique(),
        "Soma CH escalada": float(df["CARGA HORARIA ESCALADA"].fillna(0).sum()),
        "Soma CH contratada": float(df["CARGA HORARIA"].fillna(0).sum()),
        "Repouso remunerado": float(
            df["CARGA HORARIA"].fillna(0).sum()
            - df["CARGA HORARIA ESCALADA"].fillna(0).sum()
        ),
    }
    return pd.DataFrame(
        list(resumo.items()), columns=["Indicador", "Valor"]
    )


def build_por_vinculo(df):
    if len(df) == 0:
        return pd.DataFrame()
    return (
        df.groupby("VINCULO", dropna=False)
        .agg(
            QTD_ALOCACOES=("CHAVE_ID", "size"),
            QTD_SERVIDORES_UNICOS=("PESSOA_ID", lambda x: x.replace("", pd.NA).dropna().nunique()),
        )
        .reset_index()
        .sort_values("QTD_SERVIDORES_UNICOS", ascending=False)
    )


def build_por_setor(df):
    if len(df) == 0:
        return pd.DataFrame()
    return (
        df.groupby("SETOR", dropna=False)
        .agg(
            QTD_ALOCACOES=("CHAVE_ID", "size"),
            QTD_SERVIDORES_UNICOS=("PESSOA_ID", lambda x: x.replace("", pd.NA).dropna().nunique()),
            CH_ESCALADA_TOTAL=("CARGA HORARIA ESCALADA", "sum"),
        )
        .reset_index()
        .sort_values("QTD_SERVIDORES_UNICOS", ascending=False)
    )


def build_lotacao_summary(df_alerta):
    """Resumo por SETOR de um DataFrame de alerta."""
    if len(df_alerta) == 0:
        return pd.DataFrame(columns=["SETOR", "QTD_REGISTROS"])
    return (
        df_alerta.assign(_SETOR=df_alerta["SETOR"].replace("", "(SEM SETOR)"))
        .groupby("_SETOR")
        .size()
        .reset_index(name="QTD_REGISTROS")
        .rename(columns={"_SETOR": "SETOR"})
        .sort_values("QTD_REGISTROS", ascending=False)
        .reset_index(drop=True)
    )


# ==================================================================
# EXCEL FORMATADO
# ==================================================================
def style_sheet(ws, df, header_color="1F4E78"):
    header_fill = PatternFill("solid", fgColor=header_color)
    header_font = Font(bold=True, color="FFFFFF", size=11)
    thin = Side(border_style="thin", color="B7B7B7")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for idx, col_name in enumerate(df.columns, start=1):
        cell = ws.cell(row=1, column=idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(
            horizontal="center", vertical="center", wrap_text=True
        )
        cell.border = border

        max_len = len(str(col_name))
        if len(df) > 0:
            try:
                m = df[col_name].astype(str).map(len).max()
                if pd.notna(m):
                    max_len = max(max_len, int(m))
            except Exception:
                pass
        ws.column_dimensions[get_column_letter(idx)].width = min(
            max(max_len + 2, 12), 55
        )

    for row in ws.iter_rows(
        min_row=2, max_row=ws.max_row, max_col=len(df.columns)
    ):
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(vertical="center")

    if len(df) > 0:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
    ws.row_dimensions[1].height = 30


def prevent_excel_formula(value):
    """Evita que texto importado seja interpretado como fórmula no Excel."""
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def df_to_excel_bytes(sheets):
    output = BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        for name, data in sheets.items():
            safe = name[:31]
            d = data.copy()
            for c in d.select_dtypes(include=["object"]).columns:
                d[c] = d[c].map(prevent_excel_formula)
            for c in d.columns:
                if pd.api.types.is_datetime64_any_dtype(d[c]):
                    d[c] = d[c].dt.strftime("%d/%m/%Y")
            d.to_excel(writer, sheet_name=safe, index=False)
            style_sheet(writer.sheets[safe], d)
    return output.getvalue()


def df_to_excel_single_sheet(df, sheet_name="Dados"):
    return df_to_excel_bytes({sheet_name: df})


# ==================================================================
# GERAÇÃO POR LOTAÇÃO
# ==================================================================
def build_per_lotacao_files(df_alerta, rotulo, ref_date_str):
    """
    Retorna lista de dicts:
        [{ 'setor': ..., 'arquivo': ..., 'bytes': ..., 'qtd': ... }, ...]
    """
    if len(df_alerta) == 0:
        return []

    work = df_alerta.copy()
    work["_SETOR_KEY"] = work["SETOR"].replace("", "(SEM SETOR)")

    arquivos = []
    nomes_usados = set()
    for setor, grupo in work.groupby("_SETOR_KEY"):
        dados = grupo.drop(columns=["_SETOR_KEY"]).reset_index(drop=True)

        # planilha resumo
        info = pd.DataFrame({
            "Informação": [
                "Tipo de relatório",
                "Lotação / Setor",
                "Data de geração",
                "Data de referência",
                "Total de registros",
            ],
            "Valor": [
                rotulo,
                setor,
                datetime.today().strftime("%d/%m/%Y %H:%M"),
                ref_date_str,
                len(dados),
            ],
        })

        sheets = {
            "Resumo": info,
            "Registros": dados,
        }
        xls = df_to_excel_bytes(sheets)

        nome_base = sanitize_filename(f"{rotulo} - {setor}")
        nome_arquivo = f"{nome_base}.xlsx"
        contador = 2
        while nome_arquivo in nomes_usados:
            nome_arquivo = f"{nome_base}_{contador}.xlsx"
            contador += 1
        nomes_usados.add(nome_arquivo)
        arquivos.append({
            "setor": setor,
            "arquivo": nome_arquivo,
            "bytes": xls,
            "qtd": len(dados),
        })

    return sorted(arquivos, key=lambda x: x["setor"])


def build_zip_bytes(arquivos):
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        for a in arquivos:
            zf.writestr(a["arquivo"], a["bytes"])
    return buf.getvalue()


# ==================================================================
# ANÁLISE POR UNIDADE / ESCALA
# ==================================================================
def build_analise_escala(df):
    if len(df) == 0:
        return pd.DataFrame()
    out = (
        df.groupby(["CARGO", "OCUPACAO"], dropna=False)
        .agg(
            SERVIDORES_UNICOS=("PESSOA_ID", lambda x: x.replace("", pd.NA).dropna().nunique()),
            ALOCACOES=("CHAVE_ID", "size"),
            CH_ESCALADA=("CARGA HORARIA ESCALADA", "sum"),
            CH_CONTRATADA=("CARGA HORARIA", "sum"),
        ).reset_index()
    )
    out["REPOUSO_REMUNERADO"] = out["CH_CONTRATADA"] - out["CH_ESCALADA"]
    return out.sort_values(["CH_ESCALADA", "CARGO"], ascending=[False, True]).reset_index(drop=True)


def _pdf_table(data, widths=None, header=True, font_size=7):
    rows = [[str(v) if not pd.isna(v) else "" for v in row] for row in data]
    table = Table(rows, colWidths=widths, repeatRows=1 if header else 0)
    style = [
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold" if header else "Helvetica"),
        ("FONTSIZE", (0, 0), (-1, -1), font_size),
        ("LEADING", (0, 0), (-1, -1), font_size + 2),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#CBD5E1")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (-1, -1), 5),
        ("RIGHTPADDING", (0, 0), (-1, -1), 5),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]
    if header:
        style += [("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#16324F")),
                  ("TEXTCOLOR", (0, 0), (-1, 0), colors.white)]
        for r in range(1, len(rows)):
            if r % 2 == 0:
                style.append(("BACKGROUND", (0, r), (-1, r), colors.HexColor("#F1F5F9")))
    table.setStyle(TableStyle(style))
    return table


def build_pdf_report(df_filtered, unidade, escala, ref_date_str, detailed=False):
    if not HAS_REPORTLAB:
        return None
    buf = BytesIO()
    doc = SimpleDocTemplate(
        buf, pagesize=landscape(A4), rightMargin=12*mm, leftMargin=12*mm,
        topMargin=12*mm, bottomMargin=12*mm,
        title=f"Relatório {unidade} - {escala}"
    )
    styles = getSampleStyleSheet()
    title = ParagraphStyle("TitleCustom", parent=styles["Title"], fontName="Helvetica-Bold",
                           fontSize=21, leading=25, textColor=colors.HexColor("#16324F"), alignment=TA_LEFT)
    subtitle = ParagraphStyle("Sub", parent=styles["Normal"], fontSize=9, textColor=colors.HexColor("#64748B"))
    section = ParagraphStyle("Section", parent=styles["Heading2"], fontSize=12, leading=15,
                             textColor=colors.HexColor("#0F766E"), spaceBefore=10, spaceAfter=5)
    story = [Paragraph("Relatório de distribuição de carga horária", title),
             Paragraph(f"Unidade: <b>{escape(str(unidade))}</b> &nbsp;&nbsp; Escala: <b>{escape(str(escala))}</b> &nbsp;&nbsp; Referência: {escape(str(ref_date_str))}", subtitle), Spacer(1, 8)]
    pessoas = df_filtered["PESSOA_ID"].replace("", pd.NA).dropna().nunique()
    aloc = len(df_filtered)
    escalada = float(df_filtered["CARGA HORARIA ESCALADA"].fillna(0).sum())
    contratada = float(df_filtered["CARGA HORARIA"].fillna(0).sum())
    repouso = contratada - escalada
    cards = [["PESSOAS ÚNICAS", "ALOCAÇÕES", "CH ESCALADA", "CH CONTRATADA", "REPOUSO REMUNERADO"],
             [format_br_int(pessoas), format_br_int(aloc), f"{format_br_number(escalada)} h",
              f"{format_br_number(contratada)} h", f"{format_br_number(repouso, signed=True)} h"]]
    card = Table(cards, colWidths=[52*mm]*5)
    card.setStyle(TableStyle([
        ("BACKGROUND", (0,0), (-1,0), colors.HexColor("#0F766E")),
        ("TEXTCOLOR", (0,0), (-1,0), colors.white), ("FONTNAME", (0,0), (-1,0), "Helvetica-Bold"),
        ("BACKGROUND", (0,1), (-1,1), colors.HexColor("#ECFDF5")),
        ("FONTNAME", (0,1), (-1,1), "Helvetica-Bold"), ("FONTSIZE", (0,0), (-1,-1), 9),
        ("ALIGN", (0,0), (-1,-1), "CENTER"), ("BOX", (0,0), (-1,-1), 0.5, colors.HexColor("#99F6E4")),
        ("INNERGRID", (0,0), (-1,-1), 0.25, colors.HexColor("#99F6E4")),
        ("TOPPADDING", (0,0), (-1,-1), 8), ("BOTTOMPADDING", (0,0), (-1,-1), 8),
    ]))
    story += [card, Spacer(1, 10), Paragraph("Resumo por cargo e ocupação", section)]
    resumo = build_analise_escala(df_filtered)
    cols = ["CARGO", "OCUPACAO", "SERVIDORES_UNICOS", "ALOCACOES", "CH_ESCALADA", "CH_CONTRATADA", "REPOUSO_REMUNERADO"]
    labels = ["Cargo", "Ocupação/CBO", "Pessoas", "Alocações", "CH escalada", "CH contratada", "Repouso remunerado"]
    pdf_resumo = resumo[cols].fillna("").copy()
    for col in ["SERVIDORES_UNICOS", "ALOCACOES"]:
        pdf_resumo[col] = pdf_resumo[col].map(format_br_int)
    for col in ["CH_ESCALADA", "CH_CONTRATADA", "REPOUSO_REMUNERADO"]:
        pdf_resumo[col] = pdf_resumo[col].map(format_br_number)
    data = [labels] + pdf_resumo.values.tolist()
    story.append(_pdf_table(data, widths=[42*mm, 48*mm, 22*mm, 20*mm, 25*mm, 25*mm, 22*mm]))
    if detailed:
        story += [PageBreak(), Paragraph("Detalhamento dos registros", title), Spacer(1, 6)]
        detail_cols = ["SERVIDOR", "NUMFUNC", "NUMVINC", "CARGO", "OCUPACAO", "ESCALA", "CARGA HORARIA ESCALADA", "CARGA HORARIA"]
        detail_labels = ["Servidor", "Num. func.", "Vínculo", "Cargo", "Ocupação", "Escala", "CH escalada", "CH contratada"]
        details = df_filtered[detail_cols].fillna("").copy()
        for col in ["CARGA HORARIA ESCALADA", "CARGA HORARIA"]:
            details[col] = details[col].map(lambda x: format_br_number(x) if x != "" else "")
        story.append(_pdf_table([detail_labels] + details.values.tolist(), widths=[45*mm, 21*mm, 18*mm, 35*mm, 38*mm, 18*mm, 23*mm, 23*mm], font_size=6))
    generated = datetime.today().strftime("%d/%m/%Y %H:%M")
    def footer(canvas, doc_obj):
        canvas.saveState(); canvas.setFont("Helvetica", 7); canvas.setFillColor(colors.HexColor("#64748B"))
        canvas.drawString(12*mm, 7*mm, f"Gerado em {generated} | Relatório de servidores por CBO")
        canvas.drawRightString(285*mm, 7*mm, f"Página {doc_obj.page}")
        canvas.restoreState()
    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return buf.getvalue()


# ==================================================================
# UI — análise objetiva de contratos vencidos e NUMFUNC = CPF
# ==================================================================
st.markdown("""
<div class="hero">
  <div class="section-kicker" style="color:#99F6E4">Análise cadastral</div>
  <h1>Contratos vencidos</h1>
  <p>Identifique contratos temporários vencidos e registros em que o NUMFUNC é igual ao CPF.</p>
</div>
""", unsafe_allow_html=True)

with st.sidebar:
    st.header("Filtros")
    ref_date = st.date_input("Data de referência", value=datetime.today().date())
    st.caption("Contrato temporário vencido = início há mais de 2 anos da data escolhida.")

uploaded = st.file_uploader(
    "Envie a planilha de contratos",
    type=ACCEPTED_EXT,
    help="Aceita CSV, XLS, XLSX e XLSM.",
)

if uploaded is None:
    st.info("Envie uma planilha para iniciar a análise.")
    st.stop()

try:
    raw = read_any(uploaded.getvalue(), uploaded.name)
    df = clean(parse_raw(raw))
except Exception as e:
    st.error(f"Não foi possível processar `{uploaded.name}`: {e}")
    st.stop()

if len(df) == 0:
    st.warning("Nenhum registro válido encontrado na planilha.")
    st.stop()

alerta_temp = alert_contratos_vencidos(df, pd.Timestamp(ref_date))
alerta_num = alert_numfunc_cpf(df)

st.success(f"Arquivo processado: {format_br_int(len(df))} registro(s).")
k1, k2, k3 = st.columns(3)
k1.metric("Contratos vencidos", format_br_int(len(alerta_temp)))
k2.metric("NUMFUNC = CPF", format_br_int(len(alerta_num)))
k3.metric("Data de referência", pd.Timestamp(ref_date).strftime("%d/%m/%Y"))

def filtrar_por_setor(data, key):
    if len(data) == 0:
        return data
    setores = ["(Todos)"] + sorted(data["SETOR"].replace("", "(SEM SETOR)").unique().tolist())
    setor = st.selectbox("Filtrar por setor", setores, key=key)
    out = data.copy()
    out["_SETOR_FILTRO"] = out["SETOR"].replace("", "(SEM SETOR)")
    if setor != "(Todos)":
        out = out[out["_SETOR_FILTRO"] == setor]
    return out.drop(columns=["_SETOR_FILTRO"])

tab_vencidos, tab_numfunc = st.tabs(["Contratos vencidos", "NUMFUNC = CPF"])

with tab_vencidos:
    st.subheader("Contratos temporários vencidos")
    st.caption("Vínculo temporário com início há mais de 2 anos. A data de referência está na barra lateral.")
    dados_temp = filtrar_por_setor(alerta_temp, "setor_contratos_vencidos")
    if len(dados_temp) == 0:
        st.success("Nenhum contrato vencido encontrado para os filtros selecionados.")
    else:
        st.warning(f"{format_br_int(len(dados_temp))} contrato(s) vencido(s).")
        display_dataframe(dados_temp, use_container_width=True, hide_index=True, height=500)

with tab_numfunc:
    st.subheader("Registros em que NUMFUNC = CPF")
    st.caption("A comparação considera apenas os dígitos e ignora zeros à esquerda.")
    dados_num = filtrar_por_setor(alerta_num, "setor_numfunc_cpf")
    if len(dados_num) == 0:
        st.success("Nenhum registro com NUMFUNC igual ao CPF encontrado.")
    else:
        st.warning(f"{format_br_int(len(dados_num))} registro(s) encontrado(s).")
        display_dataframe(dados_num, use_container_width=True, hide_index=True, height=500)

st.markdown("---")
st.subheader("Exportar análise")
st.caption("O Excel contém somente as duas análises exibidas no painel.")
xlsx_bytes = df_to_excel_bytes({
    "Contratos Vencidos": alerta_temp,
    "NUMFUNC CPF": alerta_num,
})
st.download_button(
    "Baixar Excel da análise",
    data=xlsx_bytes,
    file_name=f"analise_contratos_vencidos_{datetime.today().strftime('%Y%m%d')}.xlsx",
    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
)
